"""Real MySQL probes in a newly-created, owned Docker tmpfs container only."""

import argparse
import json
import os
import secrets
import subprocess
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4
from unittest.mock import patch


def local_docker_context():
    context = subprocess.check_output(
        ["docker", "context", "show"], text=True, timeout=10
    ).strip()
    endpoint = json.loads(
        subprocess.check_output(
            [
                "docker",
                "context",
                "inspect",
                context,
                "--format",
                "{{json .Endpoints.docker.Host}}",
            ],
            text=True,
            timeout=10,
        )
    )
    if not endpoint.startswith(("npipe:", "unix:")) or os.environ.get("DOCKER_HOST"):
        raise RuntimeError(
            "Only an explicit local Docker Desktop/unix context without DOCKER_HOST is allowed"
        )
    return ["docker", "--context", context]


def crash_child():
    """Internal helper. All connection details come from the generated test scope."""
    from sqlalchemy.engine import make_url

    run_id = os.environ.get("HARBORRELAY_MYSQL_PROBE_RUN", "")
    url = make_url(os.environ.get("DATABASE_URL", ""))
    if (
        not run_id
        or url.database != f"harborrelay_test_{run_id}"
        or url.host != "127.0.0.1"
    ):
        raise RuntimeError("Missing isolated probe scope")
    from app.core.database import SessionLocal
    from app.core.config import get_settings
    from app.models.entities import PsychologicalReport
    from app.services.tool_queue import ToolQueueService
    from app.services.tools import ToolOrchestrationService

    with SessionLocal() as db:
        if os.environ["HARBORRELAY_PROBE_CRASH_KIND"] == "outbox":
            report = PsychologicalReport(
                user_id=int(os.environ["HARBORRELAY_PROBE_USER_ID"]),
                session_id=int(os.environ["HARBORRELAY_PROBE_SESSION_ID"]),
                content="synthetic crash probe",
                intent="CONSULT",
                emotion="NORMAL",
                emotion_score=0,
                risk_level="LOW",
                confidence=0.5,
                summary="child committed",
            )
            db.add(report)
            db.flush()
            ToolQueueService(db, get_settings()).enqueue_report(
                report.id, report.risk_level, commit=False
            )
            db.commit()
            print(report.id, flush=True)
            os._exit(73)
        report = db.get(
            PsychologicalReport, int(os.environ["HARBORRELAY_PROBE_REPORT_ID"])
        )
        db.commit = lambda: os._exit(74)
        ToolOrchestrationService(db, get_settings()).write_excel(report)


def evaluate(url, output, run_id, child_environment, threads):
    from sqlalchemy import create_engine, text
    from sqlalchemy.orm import sessionmaker
    import app.core.database as database
    from app.core.config import get_settings
    from app.models.entities import (
        UserAccount,
        ChatSession,
        ChatMessage,
        PsychologicalReport,
        ToolJob,
        ExcelRecord,
        RiskCase,
        AlertRecord,
        ToolAuditRecord,
        DeadLetterRecord,
    )

    database.engine.dispose()
    database.engine = create_engine(
        url,
        pool_size=threads + 2,
        max_overflow=0,
        pool_pre_ping=True,
        connect_args={"connect_timeout": 3, "read_timeout": 20, "write_timeout": 20},
    )
    database.SessionLocal = sessionmaker(bind=database.engine, autoflush=False)
    from app.services.tool_queue import ToolQueueService, ToolQueueWorker
    from app.services.tools import ToolOrchestrationService

    database.Base.metadata.create_all(database.engine)
    settings = get_settings()
    checks = []
    with database.SessionLocal() as db:
        user = UserAccount(
            username="synthetic-student",
            display_name="Synthetic Student",
            password_hash="unused",
        )
        db.add(user)
        db.flush()
        session = ChatSession(
            public_id=uuid4().hex, user_id=user.id, title="isolated MySQL evaluation"
        )
        db.add(session)
        db.commit()
        user_id, session_id = user.id, session.id

    def report(db, risk="HIGH"):
        row = PsychologicalReport(
            user_id=user_id,
            session_id=session_id,
            content="synthetic",
            intent="CONSULT",
            emotion="NORMAL",
            emotion_score=0,
            risk_level=risk,
            confidence=0.5,
            summary="synthetic",
        )
        db.add(row)
        db.flush()
        return row

    def check(name, function):
        started = time.perf_counter()
        try:
            detail = function()
            checks.append(
                {
                    "id": name,
                    "passed": True,
                    "details": detail,
                    "elapsed_ms": (time.perf_counter() - started) * 1000,
                }
            )
        except Exception as exc:
            checks.append(
                {
                    "id": name,
                    "passed": False,
                    "error": type(exc).__name__,
                    "elapsed_ms": (time.perf_counter() - started) * 1000,
                }
            )
        (output / "progress.json").write_text(
            json.dumps(checks, indent=2), encoding="utf-8"
        )
        print(json.dumps(checks[-1]), flush=True)

    def atomic_rollback():
        with database.SessionLocal() as db:
            row = report(db)
            report_id = row.id
            ToolQueueService(db, settings).enqueue_report(row.id, "LOW", commit=False)
            db.rollback()
        with database.SessionLocal() as db:
            assert db.get(PsychologicalReport, report_id) is None
            assert db.query(ToolJob).filter_by(report_id=report_id).count() == 0
        return {"report_and_jobs_rolled_back": True}

    check("report_outbox_atomic_rollback", atomic_rollback)

    with database.SessionLocal() as db:
        row = report(db)
        db.commit()
        shared_report_id = row.id
    shared = {}

    def concurrent_enqueue():
        rounds = []
        for _ in range(3):
            barrier = threading.Barrier(threads)

            def enqueue():
                with database.SessionLocal() as db:
                    barrier.wait(timeout=10)
                    return tuple(
                        job.id
                        for job in ToolQueueService(db, settings).enqueue_report(
                            shared_report_id, "LOW"
                        )
                    )

            with ThreadPoolExecutor(max_workers=threads) as pool:
                results = [
                    f.result(timeout=30)
                    for f in [pool.submit(enqueue) for _ in range(threads)]
                ]
            assert len(set(results)) == 1 and len(results[0]) == 3
            rounds.append(list(results[0]))
        assert all(ids == rounds[0] for ids in rounds)
        with database.SessionLocal() as db:
            assert db.query(ToolJob).filter_by(report_id=shared_report_id).count() == 3
        shared["job_ids"] = rounds[0]
        return {
            "threads": threads,
            "rounds": 3,
            "enqueue_requests": threads * 3,
            "unique_jobs": 3,
        }

    check("concurrent_same_report_enqueue", concurrent_enqueue)

    def concurrent_claim():
        submitted = []
        lock = threading.Lock()

        class Recorder:
            def submit(self, function, job_id):
                with lock:
                    submitted.append(job_id)

        workers = [ToolQueueWorker(settings) for _ in range(3)]
        barrier = threading.Barrier(3)

        def dispatch(worker):
            with patch.object(worker, "_executor_for", return_value=Recorder()):
                barrier.wait(timeout=10)
                worker._dispatch_once()

        try:
            with ThreadPoolExecutor(max_workers=3) as pool:
                for future in [pool.submit(dispatch, worker) for worker in workers]:
                    future.result(timeout=30)
            assert sorted(submitted) == sorted(shared["job_ids"])
            assert len(set(submitted)) == len(submitted) == 3
        finally:
            for worker in workers:
                worker.stop()
        return {"dispatchers": 3, "submissions": len(submitted), "duplicate_claims": 0}

    check("concurrent_conditional_claim", concurrent_claim)

    def governed_execution():
        worker = ToolQueueWorker(settings)
        try:
            with database.SessionLocal() as db:
                jobs = db.query(ToolJob).filter_by(report_id=shared_report_id).all()
                ordered = sorted(
                    jobs,
                    key=lambda job: {
                        "CASE_CREATE": 0,
                        "EXCEL_REPORT": 1,
                        "ALERT_SEND": 2,
                    }[job.kind],
                )
                job_ids = [job.id for job in ordered]
            for job_id in job_ids:
                worker._run_job(job_id)
            with database.SessionLocal() as db:
                assert all(
                    j.status == "SUCCESS"
                    for j in db.query(ToolJob).filter_by(report_id=shared_report_id)
                )
                assert (
                    db.query(ExcelRecord).filter_by(report_id=shared_report_id).count()
                    == 1
                )
                assert (
                    db.query(RiskCase).filter_by(report_id=shared_report_id).count()
                    == 1
                )
                alert = (
                    db.query(AlertRecord).filter_by(report_id=shared_report_id).one()
                )
                assert alert.status == "SUCCESS" and "deliveryMode=log" in alert.message
                assert (
                    db.query(ToolAuditRecord)
                    .filter_by(report_id=shared_report_id, status="SUCCESS")
                    .count()
                    == 3
                )
        finally:
            worker.stop()
        return {"successful_jobs": 3, "successful_audits": 3, "alert_mode": "log"}

    check("governed_jobs_and_idempotent_side_effects", governed_execution)

    env = {
        **child_environment,
        "HARBORRELAY_MYSQL_PROBE_RUN": run_id,
        "HARBORRELAY_PROBE_USER_ID": str(user_id),
        "HARBORRELAY_PROBE_SESSION_ID": str(session_id),
    }

    def committed_process_exit():
        result = subprocess.run(
            [sys.executable, "-m", "app.evaluation.mysql_runner", "--crash-child"],
            env={**env, "HARBORRELAY_PROBE_CRASH_KIND": "outbox"},
            capture_output=True,
            text=True,
            timeout=30,
        )
        assert result.returncode == 73
        report_id = int(result.stdout.strip())
        with database.SessionLocal() as db:
            assert db.get(PsychologicalReport, report_id) is not None
            job = db.query(ToolJob).filter_by(report_id=report_id).one()
            assert job.status == "PENDING"
            job.status = "RUNNING"
            job.updated_at = datetime.utcnow() - timedelta(seconds=600)
            db.commit()
            job_id = job.id
        worker = ToolQueueWorker(settings)
        try:
            worker._recover_running_jobs()
            with database.SessionLocal() as db:
                job = db.get(ToolJob, job_id)
                assert job.status == "PENDING"
                job.status = "RUNNING"
                db.commit()
            worker._run_job(job_id)
            with database.SessionLocal() as db:
                assert db.get(ToolJob, job_id).status == "SUCCESS"
        finally:
            worker.stop()
        return {
            "child_exit_code": 73,
            "committed_outbox_survived": True,
            "stale_job_recovered": True,
        }

    check("process_exit_and_stale_job_recovery", committed_process_exit)

    def excel_process_exit():
        with database.SessionLocal() as db:
            row = report(db, "LOW")
            db.commit()
            report_id = row.id
        result = subprocess.run(
            [sys.executable, "-m", "app.evaluation.mysql_runner", "--crash-child"],
            env={
                **env,
                "HARBORRELAY_PROBE_CRASH_KIND": "excel",
                "HARBORRELAY_PROBE_REPORT_ID": str(report_id),
            },
            capture_output=True,
            text=True,
            timeout=30,
        )
        assert result.returncode == 74
        with database.SessionLocal() as db:
            assert db.query(ExcelRecord).filter_by(report_id=report_id).count() == 0
            ToolOrchestrationService(db, settings).write_excel(
                db.get(PsychologicalReport, report_id)
            )
            assert db.query(ExcelRecord).filter_by(report_id=report_id).count() == 1
        from openpyxl import load_workbook

        workbook = load_workbook(settings.excel_path)
        try:
            assert (
                sum(
                    row[0] == report_id
                    for row in workbook.active.iter_rows(min_row=2, values_only=True)
                )
                == 1
            )
        finally:
            workbook.close()
        return {
            "child_exit_code": 74,
            "file_written_before_db_commit": True,
            "rows_for_report_after_retry": 1,
        }

    check("excel_process_exit_before_db_commit", excel_process_exit)

    def denied_job():
        with database.SessionLocal() as db:
            row = report(db, "LOW")
            db.commit()
            report_id = row.id
            job = ToolJob(
                report_id=report_id,
                kind="ALERT_SEND",
                status="RUNNING",
                attempts=0,
                max_attempts=3,
            )
            db.add(job)
            db.commit()
            job_id = job.id
        worker = ToolQueueWorker(settings)
        try:
            worker._run_job(job_id)
            with database.SessionLocal() as db:
                assert db.get(ToolJob, job_id).status == "DEAD"
                assert db.query(AlertRecord).filter_by(report_id=report_id).count() == 0
                assert (
                    db.query(ToolAuditRecord)
                    .filter_by(job_id=job_id, allowed=False, status="BLOCKED")
                    .count()
                    == 1
                )
        finally:
            worker.stop()
        return {"unauthorized_alerts": 0, "refusal_audited": True}

    check("queue_policy_refusal_on_real_mysql", denied_job)

    def dependency_dead():
        with database.SessionLocal() as db:
            row = report(db)
            db.commit()
            report_id = row.id
            jobs = ToolQueueService(db, settings).enqueue_report(report_id, "HIGH")
            case = next(j for j in jobs if j.kind == "CASE_CREATE")
            case.status = "DEAD"
            alert = next(j for j in jobs if j.kind == "ALERT_SEND")
            alert.status = "RUNNING"
            db.commit()
            job_id = alert.id
        worker = ToolQueueWorker(settings)
        try:
            worker._run_job(job_id)
            with database.SessionLocal() as db:
                assert db.get(ToolJob, job_id).status == "DEAD"
                assert db.query(DeadLetterRecord).filter_by(job_id=job_id).count() == 1
        finally:
            worker.stop()
        return {"dependent_job_dead_lettered": True}

    check("dead_dependency_converges", dependency_dead)

    def message_window_commit_order():
        from app.harness.runner import InMemoryShortTermMemoryStore
        from app.services.memory import load_committed_history

        memory = InMemoryShortTermMemoryStore(settings)
        memory.reset()
        with database.SessionLocal() as first, database.SessionLocal() as second:
            early = ChatMessage(
                user_id=user_id,
                session_id=session_id,
                role="USER",
                content="synthetic allocated first",
            )
            first.add(early)
            first.flush()
            later = ChatMessage(
                user_id=user_id,
                session_id=session_id,
                role="USER",
                content="synthetic committed first",
            )
            second.add(later)
            second.commit()
            late_id = later.id
            memory.append("synthetic-window", "user", later.content, message_id=late_id)
            with database.SessionLocal() as reader:
                messages, source = load_committed_history(
                    reader, memory, session_id, "synthetic-window", 4
                )
                assert source == "redis_verified" and len(messages) == 1
            first.commit()
            assert early.id < late_id
        with database.SessionLocal() as reader:
            messages, source = load_committed_history(
                reader, memory, session_id, "synthetic-window", 4
            )
            assert source == "mysql_seeded" and len(messages) == 2
            assert messages[0].content == "synthetic allocated first"
            assert (
                load_committed_history(
                    reader, memory, session_id, "synthetic-window", 4
                )[1]
                == "redis_verified"
            )
        memory.reset()
        return {
            "two_transactions_reverse_commit_order": True,
            "late_lower_id_detected": True,
            "bounded_complete_window_validated": True,
            "cache_adapter": "in_memory; real Redis in infrastructure probe",
        }

    check(
        "committed_message_window_handles_reverse_commit_order",
        message_window_commit_order,
    )
    with database.engine.connect() as connection:
        version = connection.execute(text("SELECT VERSION()")).scalar()
        isolation = connection.execute(text("SELECT @@transaction_isolation")).scalar()
    database.engine.dispose()
    return {
        "checks": checks,
        "passed": all(c["passed"] for c in checks),
        "server_version": version,
        "transaction_isolation": isolation,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", default="target/mysql-validation")
    parser.add_argument("--threads", type=int, default=8, choices=range(2, 13))
    parser.add_argument("--crash-child", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if args.crash_child:
        crash_child()
        return 1
    output = Path(args.output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    run_id = uuid4().hex[:12]
    password = secrets.token_hex(20)
    container = None
    docker = None
    report = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "run_id": run_id,
        "isolation": "new owned local Docker container; tmpfs data; no existing DB/volume access",
        "smtp_mode": "log",
        "passed": False,
        "cleanup_succeeded": False,
    }
    try:
        docker = local_docker_context()
        image = json.loads(
            subprocess.check_output(
                [*docker, "image", "inspect", "mysql:8.0", "--format", "{{json .Id}}"],
                text=True,
                timeout=10,
            )
        )
        environment = {
            **os.environ,
            "MYSQL_ROOT_PASSWORD": password,
            "MYSQL_DATABASE": f"harborrelay_test_{run_id}",
        }
        container = subprocess.check_output(
            [
                *docker,
                "run",
                "--pull=never",
                "--rm",
                "-d",
                "--name",
                f"harborrelay-probe-{run_id}",
                "--label",
                f"harborrelay.validation-run={run_id}",
                "--label",
                "harborrelay.synthetic=true",
                "--memory",
                "1g",
                "--cpus",
                "1",
                "--tmpfs",
                "/var/lib/mysql:rw,size=512m",
                "-p",
                "127.0.0.1::3306",
                "--env",
                "MYSQL_ROOT_PASSWORD",
                "--env",
                "MYSQL_DATABASE",
                "mysql:8.0",
                "--skip-log-bin",
                "--innodb-buffer-pool-size=64M",
            ],
            env=environment,
            text=True,
            timeout=30,
        ).strip()
        endpoint = subprocess.check_output(
            [*docker, "port", container, "3306/tcp"], text=True, timeout=10
        ).strip()
        if not endpoint.startswith("127.0.0.1:"):
            raise RuntimeError("Unexpected Docker port binding")
        port = int(endpoint.rsplit(":", 1)[1])
        url = f"mysql+pymysql://root:{password}@127.0.0.1:{port}/harborrelay_test_{run_id}?charset=utf8mb4"
        import pymysql

        deadline = time.monotonic() + 90
        while True:
            try:
                connection = pymysql.connect(
                    host="127.0.0.1",
                    port=port,
                    user="root",
                    password=password,
                    database=f"harborrelay_test_{run_id}",
                    connect_timeout=2,
                )
                connection.close()
                break
            except pymysql.MySQLError:
                if time.monotonic() > deadline:
                    raise TimeoutError("Isolated MySQL did not become ready")
                time.sleep(1)
        os.environ.update(
            DATABASE_URL=url,
            AI_PROVIDER="mock",
            KNOWLEDGE_VECTOR_ENABLED="false",
            ALERT_EMAIL_DELIVERY_MODE="log",
            EXCEL_PATH=str(output / f"synthetic-ledger-{run_id}.xlsx"),
            TOOL_QUEUE_ENABLED="false",
        )
        report.update(evaluate(url, output, run_id, dict(os.environ), args.threads))
        report["image_id"] = image
    except Exception as exc:
        report["error"] = type(exc).__name__
        report["detail"] = str(exc).replace(password, "[generated-test-secret]")[:1000]
    finally:
        if container is not None and docker is not None:
            marker = subprocess.check_output(
                [
                    *docker,
                    "inspect",
                    container,
                    "--format",
                    '{{index .Config.Labels "harborrelay.validation-run"}}',
                ],
                text=True,
                timeout=10,
            ).strip()
            if marker != run_id:
                raise RuntimeError(
                    "Refusing cleanup of a container without this run's ownership label"
                )
            subprocess.run(
                [*docker, "stop", "--time", "10", container],
                check=True,
                capture_output=True,
                timeout=25,
            )
            report["cleanup_succeeded"] = True
        (output / "report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(
        json.dumps(
            {key: value for key, value in report.items() if key != "checks"},
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if report["passed"] and report["cleanup_succeeded"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
