import asyncio
import base64
import hashlib
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.requests import Request

from app.agents.factory import create_agent_runtime
from app.agents.harness import MindBridgeAgentHarness
from app.core.config import Settings
from app.core.database import Base
from app.core.enums import EmotionLabel, RiskLevel
from app.core.security import current_user, verify_password
from app.harness.runner import InMemoryShortTermMemoryStore
from app.models.entities import (
    AlertRecord,
    ChatSession,
    PsychologicalReport,
    RiskCase,
    ToolAuditRecord,
    ToolJob,
    UserAccount,
    DeadLetterRecord,
)
from app.schemas.dtos import ChatRequest
from app.services.assessment import PsychologyAssessment
from app.services.chat import ChatService
from app.services.report import ReportService
from app.services.safety import FinalResponseReview, HIGH_RISK_RESPONSE
from app.services.tool_queue import ToolQueueService, ToolQueueWorker
from app.services.tools import ToolOrchestrationService


class SafetyBoundariesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.engine = create_engine(
            "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
        )
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.settings = Settings(
            _env_file=None,
            ai_provider="mock",
            knowledge_vector_enabled=False,
            tool_queue_enabled=True,
            alert_email_delivery_mode="log",
            excel_path=str(Path(self.temp.name) / "ledger.xlsx"),
        )
        self.user = UserAccount(
            username="student", display_name="Student", password_hash="unused"
        )
        self.staff = UserAccount(
            username="counselor",
            display_name="Counselor",
            password_hash="unused",
            roles_csv="ROLE_COUNSELOR",
        )
        self.other = UserAccount(
            username="other",
            display_name="Other",
            password_hash="unused",
            roles_csv="ROLE_COUNSELOR",
        )
        self.db.add_all([self.user, self.staff, self.other])
        self.db.commit()
        self.patches = [
            patch(
                "app.agents.runtime.RedisShortTermMemoryStore",
                InMemoryShortTermMemoryStore,
            ),
            patch(
                "app.agents.harness.RedisShortTermMemoryStore",
                InMemoryShortTermMemoryStore,
            ),
        ]
        for value in self.patches:
            value.start()
        InMemoryShortTermMemoryStore.reset()

    def tearDown(self):
        for value in reversed(self.patches):
            value.stop()
        self.db.close()
        self.engine.dispose()
        self.temp.cleanup()

    def runtime(self, framework="custom"):
        self.settings.agent_framework = framework
        session = ChatSession(
            public_id=f"session-{framework}", user_id=self.user.id, title="synthetic"
        )
        self.db.add(session)
        self.db.commit()
        return create_agent_runtime(self.db, self.settings), session

    def report(self, risk="HIGH"):
        _, session = self.runtime()
        report = PsychologicalReport(
            user_id=self.user.id,
            session_id=session.id,
            content="synthetic",
            intent="CONSULT",
            emotion="NORMAL",
            emotion_score=0,
            risk_level=risk,
            confidence=0.8,
            summary="synthetic",
        )
        self.db.add(report)
        self.db.commit()
        return report

    def test_all_intents_force_risk_assessment_in_both_runtimes(self):
        for framework in ["custom", "langgraph"]:
            with self.subTest(framework=framework):
                runtime, session = self.runtime(framework)
                with patch.object(
                    runtime.assessment, "assess", wraps=runtime.assessment.assess
                ) as assess:
                    result = runtime.run(
                        self.user, session, "解释 Python 函数", "解释 Python 函数"
                    )
                assess.assert_called_once()
                names = [step.agent for step in result.steps]
                self.assertLess(
                    names.index("RiskGuardianAgent"), names.index("SupervisorAgent")
                )
                self.assertEqual(result.intent.value, "CHAT")

    def test_assessment_override_cannot_be_lowered_by_chat_classifier(self):
        runtime, session = self.runtime()
        assessment = PsychologyAssessment(
            EmotionLabel.HIGH_RISK, 4, RiskLevel.HIGH, 0.9, "synthetic"
        )
        with (
            patch.object(runtime.assessment, "assess", return_value=assessment),
            patch.object(runtime, "_classify", return_value="CHAT"),
        ):
            result = runtime.run(
                self.user, session, "解释 Python 函数", "解释 Python 函数"
            )
        self.assertEqual(result.risk_level, RiskLevel.HIGH)
        self.assertEqual(result.intent.value, "RISK")

    def test_mixed_general_question_and_risk_is_high_without_rag(self):
        runtime, session = self.runtime()
        with patch.object(
            runtime.knowledge,
            "retrieve",
            side_effect=AssertionError("must not delay safety"),
        ):
            result = runtime.run(
                self.user,
                session,
                "帮我写 Python，但我不想活了",
                "帮我写 Python，但我不想活了",
            )
        self.assertEqual(result.risk_level, RiskLevel.HIGH)

    def test_assessment_exception_is_conservative(self):
        runtime, session = self.runtime()
        with patch.object(runtime.assessment, "assess", side_effect=TimeoutError):
            result = runtime.run(
                self.user, session, "解释 Python 函数", "解释 Python 函数"
            )
        self.assertEqual(result.risk_level, RiskLevel.MEDIUM)
        self.assertEqual(result.intent.value, "CONSULT")

    def test_disconnect_after_meta_keeps_committed_outbox(self):
        async def run():
            stream = ChatService(self.db, self.settings).stream_chat(
                self.user, ChatRequest(message="我不想活了")
            )
            self.assertIn("event: meta", await anext(stream))
            await stream.aclose()

        asyncio.run(run())
        self.assertEqual(self.db.query(PsychologicalReport).count(), 1)
        self.assertEqual(self.db.query(ToolJob).count(), 3)

    def test_high_risk_final_template_is_checked_and_never_calls_generation(self):
        service = ChatService(self.db, self.settings)
        with patch.object(
            service.ai, "stream", side_effect=AssertionError("unreviewed text")
        ):
            events = asyncio.run(self.collect(service, "我不想活了"))
        self.assertIn(HIGH_RISK_RESPONSE, "".join(events))
        self.assertEqual(self.db.query(ToolJob).count(), 3)

    def test_review_refusal_never_emits_candidate_or_fallback(self):
        with patch(
            "app.services.chat.review_final_response",
            return_value=FinalResponseReview(False, "rejected", "safety_template"),
        ):
            events = asyncio.run(
                self.collect(ChatService(self.db, self.settings), "我不想活了")
            )
        self.assertFalse(any("event: token" in event for event in events))
        self.assertTrue(any("event: error" in event for event in events))
        self.assertEqual(self.db.query(ToolJob).count(), 3)

    def test_model_stream_failure_preserves_outbox(self):
        service = ChatService(self.db, self.settings)

        async def fail(messages):
            raise TimeoutError("synthetic")
            yield "unreachable"

        with patch.object(service.ai, "stream", fail):
            events = asyncio.run(self.collect(service, "最近压力很大，焦虑失眠"))
        self.assertTrue(any("event: error" in event for event in events))
        self.assertGreater(self.db.query(ToolJob).count(), 0)

    async def collect(self, service, message):
        return [
            event
            async for event in service.stream_chat(
                self.user, ChatRequest(message=message)
            )
        ]

    def test_report_and_outbox_rollback_together(self):
        with patch.object(
            ToolQueueService, "enqueue_report", side_effect=RuntimeError("synthetic")
        ):
            with self.assertRaises(RuntimeError):
                MindBridgeAgentHarness(self.db, self.settings).run(
                    self.user, ChatRequest(message="我不想活了")
                )
        self.db.rollback()
        self.assertEqual(self.db.query(PsychologicalReport).count(), 0)
        self.assertEqual(self.db.query(ToolJob).count(), 0)

    def test_enqueue_uses_persisted_risk_and_is_idempotent(self):
        report = self.report()
        queue = ToolQueueService(self.db, self.settings)
        first = queue.enqueue_report(report.id, "LOW")
        second = queue.enqueue_report(report.id, "HIGH")
        self.assertEqual([j.id for j in first], [j.id for j in second])
        self.assertEqual(len(first), 3)

    def test_queue_policy_refusal_and_success_are_audited(self):
        report = self.report("LOW")
        worker = ToolQueueWorker(self.settings)
        denied = ToolJob(
            report_id=report.id,
            kind="ALERT_SEND",
            status="RUNNING",
            attempts=1,
            max_attempts=3,
        )
        allowed = ToolJob(
            report_id=report.id,
            kind="EXCEL_REPORT",
            status="RUNNING",
            attempts=1,
            max_attempts=3,
        )
        self.db.add_all([denied, allowed])
        self.db.commit()
        try:
            with self.assertRaises(PermissionError):
                worker._execute(self.db, denied)
            worker._execute(self.db, allowed)
            self.assertEqual(self.db.query(AlertRecord).count(), 0)
            self.assertEqual(
                {row.status for row in self.db.query(ToolAuditRecord)},
                {"BLOCKED", "SUCCESS"},
            )
        finally:
            worker.stop()

    def test_tool_failure_retries_then_succeeds_without_real_email(self):
        report = self.report()
        job = ToolQueueService(self.db, self.settings).enqueue_report(
            report.id, report.risk_level
        )[0]
        job.status = "RUNNING"
        self.db.commit()
        worker = ToolQueueWorker(self.settings)
        with patch(
            "app.services.tool_queue.SessionLocal", sessionmaker(bind=self.engine)
        ):
            with patch.object(
                ToolOrchestrationService,
                "write_excel",
                side_effect=TimeoutError("synthetic"),
            ):
                worker._run_job(job.id)
            self.db.expire_all()
            job = self.db.get(ToolJob, job.id)
            self.assertEqual(job.status, "PENDING")
            job.status = "RUNNING"
            self.db.commit()
            worker._run_job(job.id)
            self.db.expire_all()
            self.assertEqual(self.db.get(ToolJob, job.id).status, "SUCCESS")
        worker.stop()

    def test_forged_actor_student_and_other_owner_are_denied(self):
        case = RiskCase(
            report_id=1,
            risk_level="HIGH",
            status="OPEN",
            owner=self.staff.username,
            summary="synthetic",
        )
        self.db.add(case)
        self.db.commit()
        tools = ToolOrchestrationService(self.db, self.settings)
        for actor in ["admin", self.user, self.other]:
            with self.subTest(actor=str(actor)):
                with self.assertRaises(PermissionError):
                    tools.acknowledge_case(case.id, actor)
                with self.assertRaises(PermissionError):
                    ReportService(self.db).case_notes(case.id, actor)
        self.assertEqual(
            self.db.query(ToolAuditRecord).filter_by(allowed=False).count(), 3
        )

    def test_mcp_mutations_refuse_untrusted_actor(self):
        from app.mcp_tools.server import (
            mindbridge_alert_ack,
            mindbridge_case_note_add,
            mindbridge_excel_report,
        )

        for tool, args in [
            (mindbridge_alert_ack, (1, "admin")),
            (mindbridge_case_note_add, (1, "admin", "note")),
            (mindbridge_excel_report, (1,)),
        ]:
            with self.assertRaises(PermissionError):
                tool(*args)

    def test_password_compatibility_and_successful_login_upgrade(self):
        legacy = hashlib.sha256(b"synthetic-password").hexdigest()
        self.user.password_hash = legacy
        self.db.commit()
        self.assertTrue(verify_password("synthetic-password", legacy))
        self.assertFalse(verify_password("wrong", legacy))
        authorization = base64.b64encode(b"student:synthetic-password")
        request = Request(
            {"type": "http", "headers": [(b"authorization", b"Basic " + authorization)]}
        )
        principal = current_user(request, self.db)
        self.assertTrue(principal.password_hash.startswith("$argon2id$"))
        self.assertTrue(verify_password("synthetic-password", principal.password_hash))
        self.assertFalse(verify_password("wrong", principal.password_hash))
        self.assertFalse(verify_password("wrong", "malformed"))

    def test_unicode_or_whitespace_obfuscation_does_not_bypass_risk_rule(self):
        from app.services.ai import has_high_risk_signal

        for text in ["我想自\u200b杀", "我想自 杀", "I want to ＫＩＬＬ myself"]:
            self.assertTrue(has_high_risk_signal(text))

    def test_review_timeout_never_emits_unapproved_text(self):
        with patch("app.services.chat.review_final_response", side_effect=TimeoutError):
            events = asyncio.run(
                self.collect(ChatService(self.db, self.settings), "我不想活了")
            )
        self.assertFalse(any("event: token" in event for event in events))
        self.assertTrue(any("event: error" in event for event in events))
        self.assertEqual(self.db.query(ToolJob).count(), 3)

    def test_excel_retry_after_file_written_database_rollback_does_not_duplicate_row(
        self,
    ):
        from openpyxl import load_workbook

        report = self.report()
        service = ToolOrchestrationService(self.db, self.settings)
        with patch.object(
            self.db, "commit", side_effect=RuntimeError("synthetic commit failure")
        ):
            with self.assertRaises(RuntimeError):
                service.write_excel(report)
        self.db.rollback()
        service.write_excel(report)
        workbook = load_workbook(self.settings.excel_path)
        try:
            self.assertEqual(workbook.active.max_row, 2)
        finally:
            workbook.close()

    def test_dead_dependency_converges_to_dead_letter(self):
        report = self.report()
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(
            report.id, report.risk_level
        )
        case = next(j for j in jobs if j.kind == "CASE_CREATE")
        alert = next(j for j in jobs if j.kind == "ALERT_SEND")
        case.status = "DEAD"
        alert.status = "RUNNING"
        self.db.commit()
        worker = ToolQueueWorker(self.settings)
        with patch(
            "app.services.tool_queue.SessionLocal", sessionmaker(bind=self.engine)
        ):
            worker._run_job(alert.id)
        self.db.expire_all()
        self.assertEqual(self.db.get(ToolJob, alert.id).status, "DEAD")
        self.assertEqual(
            self.db.query(DeadLetterRecord).filter_by(job_id=alert.id).count(), 1
        )
        worker.stop()

    def test_recovery_does_not_reclaim_fresh_running_job(self):
        report = self.report()
        stale = ToolJob(
            report_id=report.id,
            kind="EXCEL_REPORT",
            status="RUNNING",
            updated_at=datetime.utcnow() - timedelta(seconds=600),
        )
        fresh = ToolJob(
            report_id=report.id,
            kind="CASE_CREATE",
            status="RUNNING",
            updated_at=datetime.utcnow(),
        )
        self.db.add_all([stale, fresh])
        self.db.commit()
        worker = ToolQueueWorker(self.settings)
        with patch(
            "app.services.tool_queue.SessionLocal", sessionmaker(bind=self.engine)
        ):
            worker._recover_running_jobs()
        self.db.expire_all()
        self.assertEqual(self.db.get(ToolJob, stale.id).status, "PENDING")
        self.assertEqual(self.db.get(ToolJob, fresh.id).status, "RUNNING")
        worker.stop()

    def test_dead_job_is_not_silently_recreated_by_enqueue(self):
        report = self.report("LOW")
        queue = ToolQueueService(self.db, self.settings)
        job = queue.enqueue_report(report.id, "LOW")[0]
        job.status = "DEAD"
        self.db.commit()
        self.assertEqual(queue.enqueue_report(report.id, "LOW")[0].id, job.id)
        self.assertEqual(self.db.query(ToolJob).count(), 1)

    def test_disconnect_is_recorded_in_trace(self):
        import json
        from app.models.entities import AgentRunTrace

        async def run():
            stream = ChatService(self.db, self.settings).stream_chat(
                self.user, ChatRequest(message="我不想活了")
            )
            await anext(stream)
            await stream.aclose()

        asyncio.run(run())
        trace = self.db.query(AgentRunTrace).one()
        self.assertEqual(
            json.loads(trace.assessment_json)["response_stop_reason"],
            "client_disconnected",
        )

    def test_optional_rag_skills_and_cache_cannot_block_high_risk_outbox(self):
        with (
            patch(
                "app.agents.runtime.KnowledgeService",
                side_effect=RuntimeError("optional RAG unavailable"),
            ),
            patch(
                "app.agents.runtime.RedisShortTermMemoryStore",
                side_effect=RuntimeError("cache unavailable"),
            ),
            patch(
                "app.agents.harness.RedisShortTermMemoryStore",
                side_effect=RuntimeError("cache unavailable"),
            ),
            patch(
                "app.agents.runtime.MindBridgeSkillLibrary.response_skill_context",
                side_effect=RuntimeError("skills unavailable"),
            ),
        ):
            events = asyncio.run(
                self.collect(ChatService(self.db, self.settings), "我不想活了")
            )
        self.assertIn(HIGH_RISK_RESPONSE, "".join(events))
        self.assertEqual(self.db.query(ToolJob).count(), 3)

    def test_cache_initialization_failure_cannot_skip_normal_risk_assessment(self):
        runtime, session = self.runtime()
        with patch(
            "app.agents.runtime.RedisShortTermMemoryStore",
            side_effect=RuntimeError("cache unavailable"),
        ):
            result = runtime.run(
                self.user, session, "解释 Python 函数", "解释 Python 函数"
            )
        self.assertIn("RiskGuardianAgent", [step.agent for step in result.steps])

    def test_normal_sse_yields_provider_chunks_before_generation_finishes(self):
        service = ChatService(self.db, self.settings)
        produced = []

        async def provider(messages):
            produced.append("first")
            yield "first"
            produced.append("second")
            yield "second"

        async def run():
            stream = service.stream_chat(self.user, ChatRequest(message="解释 Python 函数"))
            self.assertIn("event: meta", await anext(stream))
            self.assertEqual(produced, [])
            self.assertIn('"content": "first"', await anext(stream))
            self.assertEqual(produced, ["first"])
            self.assertIn('"content": "second"', await anext(stream))
            self.assertEqual(produced, ["first", "second"])
            await stream.aclose()

        with patch.object(service.ai, "stream", provider):
            asyncio.run(run())


if __name__ == "__main__":
    unittest.main()
